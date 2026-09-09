import streamlit as st

st.title(":blue[MATERIAL STOK SPAREPART]")
st.title("Aplikasi Material HUT Purwokerto :blue[] :sunglasses:")
st.title("Dashboard", icon=":material/dashboard:")

```python
import streamlit as st
import sqlite3
import pandas as pd
import io
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =========================================================
# CONFIG
# =========================================================

st.set_page_config(
    page_title="Sistem Stok Gudang",
    page_icon="📦",
    layout="wide"
)

DB_NAME = "gudang.db"
UPLOAD_DIR = "uploads"

os.makedirs(UPLOAD_DIR, exist_ok=True)


# =========================================================
# DATABASE
# =========================================================

def get_connection():
    return sqlite3.connect(DB_NAME)


def init_database():

    conn = get_connection()
    cursor = conn.cursor()

    # Master barang
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS barang (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kode_barang TEXT UNIQUE NOT NULL,
            nama_barang TEXT NOT NULL,
            kategori TEXT,
            satuan TEXT,
            stok_minimum INTEGER DEFAULT 0,
            lokasi TEXT,
            created_at TEXT
        )
    """)

    # Transaksi
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS transaksi (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nomor_dokumen TEXT,
            tanggal TEXT,
            jenis TEXT,
            kode_barang TEXT,
            nama_barang TEXT,
            jumlah INTEGER,
            satuan TEXT,
            supplier_customer TEXT,
            keterangan TEXT,
            foto TEXT,
            created_at TEXT
        )
    """)

    conn.commit()
    conn.close()


init_database()


# =========================================================
# FUNCTIONS
# =========================================================

def get_barang():

    conn = get_connection()

    df = pd.read_sql_query(
        "SELECT * FROM barang ORDER BY nama_barang",
        conn
    )

    conn.close()

    return df


def get_transaksi():

    conn = get_connection()

    df = pd.read_sql_query(
        "SELECT * FROM transaksi ORDER BY tanggal DESC, id DESC",
        conn
    )

    conn.close()

    return df


def tambah_barang(
    kode,
    nama,
    kategori,
    satuan,
    stok_minimum,
    lokasi
):

    conn = get_connection()

    try:

        conn.execute("""
            INSERT INTO barang
            (
                kode_barang,
                nama_barang,
                kategori,
                satuan,
                stok_minimum,
                lokasi,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            kode,
            nama,
            kategori,
            satuan,
            stok_minimum,
            lokasi,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))

        conn.commit()

        return True, "Barang berhasil ditambahkan."

    except sqlite3.IntegrityError:

        return False, "Kode barang sudah digunakan."

    finally:

        conn.close()


def tambah_transaksi(
    nomor_dokumen,
    tanggal,
    jenis,
    kode_barang,
    nama_barang,
    jumlah,
    satuan,
    supplier_customer,
    keterangan,
    foto
):

    conn = get_connection()

    conn.execute("""
        INSERT INTO transaksi
        (
            nomor_dokumen,
            tanggal,
            jenis,
            kode_barang,
            nama_barang,
            jumlah,
            satuan,
            supplier_customer,
            keterangan,
            foto,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        nomor_dokumen,
        tanggal,
        jenis,
        kode_barang,
        nama_barang,
        jumlah,
        satuan,
        supplier_customer,
        keterangan,
        foto,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))

    conn.commit()
    conn.close()


def hitung_stok():

    barang = get_barang()
    transaksi = get_transaksi()

    if barang.empty:
        return pd.DataFrame()

    if transaksi.empty:

        barang["stok_masuk"] = 0
        barang["stok_keluar"] = 0
        barang["stok"] = 0

    else:

        masuk = (
            transaksi[transaksi["jenis"] == "Barang Masuk"]
            .groupby("kode_barang")["jumlah"]
            .sum()
        )

        keluar = (
            transaksi[transaksi["jenis"] == "Barang Keluar"]
            .groupby("kode_barang")["jumlah"]
            .sum()
        )

        barang["stok_masuk"] = (
            barang["kode_barang"]
            .map(masuk)
            .fillna(0)
            .astype(int)
        )

        barang["stok_keluar"] = (
            barang["kode_barang"]
            .map(keluar)
            .fillna(0)
            .astype(int)
        )

        barang["stok"] = (
            barang["stok_masuk"]
            - barang["stok_keluar"]
        )

    barang["status"] = barang.apply(
        lambda x:
            "⚠️ STOK MINIMUM"
            if x["stok"] <= x["stok_minimum"]
            else "OK",
        axis=1
    )

    return barang


def export_excel():

    stok = hitung_stok()
    transaksi = get_transaksi()
    barang = get_barang()

    output = io.BytesIO()

    wb = Workbook()

    # =====================================================
    # SHEET STOK
    # =====================================================

    ws = wb.active
    ws.title = "Stok Barang"

    headers = [
        "Kode Barang",
        "Nama Barang",
        "Kategori",
        "Satuan",
        "Lokasi",
        "Stok Minimum",
        "Stok Masuk",
        "Stok Keluar",
        "Stok Akhir",
        "Status"
    ]

    ws.append(headers)

    for cell in ws[1]:

        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    for _, row in stok.iterrows():

        ws.append([
            row["kode_barang"],
            row["nama_barang"],
            row["kategori"],
            row["satuan"],
            row["lokasi"],
            row["stok_minimum"],
            row["stok_masuk"],
            row["stok_keluar"],
            row["stok"],
            row["status"]
        ])

    # =====================================================
    # SHEET TRANSAKSI
    # =====================================================

    ws2 = wb.create_sheet("Riwayat Transaksi")

    headers2 = [
        "No",
        "Nomor Dokumen",
        "Tanggal",
        "Jenis",
        "Kode Barang",
        "Nama Barang",
        "Jumlah",
        "Satuan",
        "Supplier / Customer",
        "Keterangan",
        "Foto"
    ]

    ws2.append(headers2)

    for cell in ws2[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="70AD47"
        )

    for _, row in transaksi.iterrows():

        ws2.append([
            row["id"],
            row["nomor_dokumen"],
            row["tanggal"],
            row["jenis"],
            row["kode_barang"],
            row["nama_barang"],
            row["jumlah"],
            row["satuan"],
            row["supplier_customer"],
            row["keterangan"],
            row["foto"]
        ])

    # =====================================================
    # SHEET MASTER
    # =====================================================

    ws3 = wb.create_sheet("Master Barang")

    headers3 = [
        "ID",
        "Kode Barang",
        "Nama Barang",
        "Kategori",
        "Satuan",
        "Stok Minimum",
        "Lokasi"
    ]

    ws3.append(headers3)

    for cell in ws3[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="ED7D31"
        )

    for _, row in barang.iterrows():

        ws3.append([
            row["id"],
            row["kode_barang"],
            row["nama_barang"],
            row["kategori"],
            row["satuan"],
            row["stok_minimum"],
            row["lokasi"]
        ])

    # =====================================================
    # AUTO WIDTH
    # =====================================================

    for sheet in wb.worksheets:

        for column in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                try:
                    if cell.value:
                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )
                except:
                    pass

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                40
            )

    wb.save(output)

    output.seek(0)

    return output


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title("📦 Gudang")

menu = st.sidebar.radio(
    "Menu",
    [
        "Dashboard",
        "Master Barang",
        "Barang Masuk",
        "Barang Keluar",
        "Riwayat Transaksi",
        "Laporan Excel"
    ]
)


# =========================================================
# DASHBOARD
# =========================================================

if menu == "Dashboard":

    st.title("📊 Dashboard Stok Gudang")

    stok = hitung_stok()
    transaksi = get_transaksi()

    if stok.empty:

        st.info(
            "Belum ada master barang. "
            "Silakan tambahkan barang terlebih dahulu."
        )

    else:

        total_barang = len(stok)

        total_stok = stok["stok"].sum()

        stok_minimum = len(
            stok[
                stok["stok"]
                <= stok["stok_minimum"]
            ]
        )

        total_masuk = stok["stok_masuk"].sum()

        total_keluar = stok["stok_keluar"].sum()

        col1, col2, col3, col4, col5 = st.columns(5)

        col1.metric(
            "Jenis Barang",
            total_barang
        )

        col2.metric(
            "Total Stok",
            total_stok
        )

        col3.metric(
            "Barang Masuk",
            total_masuk
        )

        col4.metric(
            "Barang Keluar",
            total_keluar
        )

        col5.metric(
            "Stok Minimum",
            stok_minimum
        )

        st.divider()

        st.subheader("📦 Kondisi Stok")

        tampil = stok[
            [
                "kode_barang",
                "nama_barang",
                "kategori",
                "satuan",
                "lokasi",
                "stok_minimum",
                "stok_masuk",
                "stok_keluar",
                "stok",
                "status"
            ]
        ]

        tampil.columns = [
            "Kode",
            "Nama Barang",
            "Kategori",
            "Satuan",
            "Lokasi",
            "Stok Minimum",
            "Masuk",
            "Keluar",
            "Stok Akhir",
            "Status"
        ]

        st.dataframe(
            tampil,
            use_container_width=True,
            hide_index=True
        )


# =========================================================
# MASTER BARANG
# =========================================================

elif menu == "Master Barang":

    st.title("📦 Master Barang")

    with st.expander(
        "➕ Tambah Barang",
        expanded=True
    ):

        col1, col2 = st.columns(2)

        with col1:

            kode = st.text_input(
                "Kode Barang *"
            )

            nama = st.text_input(
                "Nama Barang *"
            )

            kategori = st.text_input(
                "Kategori"
            )

        with col2:

            satuan = st.selectbox(
                "Satuan",
                [
                    "PCS",
                    "UNIT",
                    "BOX",
                    "KG",
                    "LITER",
                    "METER",
                    "SET"
                ]
            )

            stok_minimum = st.number_input(
                "Stok Minimum",
                min_value=0,
                value=0
            )

            lokasi = st.text_input(
                "Lokasi/Rak"
            )

        if st.button(
            "💾 Simpan Barang",
            type="primary"
        ):

            if not kode or not nama:

                st.error(
                    "Kode dan nama barang wajib diisi."
                )

            else:

                success, message = tambah_barang(
                    kode,
                    nama,
                    kategori,
                    satuan,
                    stok_minimum,
                    lokasi
                )

                if success:

                    st.success(message)
                    st.rerun()

                else:

                    st.error(message)

    st.divider()

    st.subheader("Daftar Master Barang")

    df = get_barang()

    st.dataframe(
        df,
        use_container_width=True,
        hide_index=True
    )


# =========================================================
# BARANG MASUK / KELUAR
# =========================================================

elif menu in [
    "Barang Masuk",
    "Barang Keluar"
]:

    jenis = menu

    st.title(
        "📥 " if jenis == "Barang Masuk"
        else "📤 "
        + jenis
    )

    barang = get_barang()

    if barang.empty:

        st.warning(
            "Master barang masih kosong. "
            "Tambahkan barang terlebih dahulu."
        )

    else:

        kode_list = barang[
            "kode_barang"
        ].tolist()

        with st.form("form_transaksi"):

            col1, col2 = st.columns(2)

            with col1:

                nomor_dokumen = st.text_input(
                    "Nomor Dokumen"
                )

                tanggal = st.date_input(
                    "Tanggal",
                    value=datetime.today()
                )

                kode_barang = st.selectbox(
                    "Kode Barang",
                    kode_list
                )

                selected = barang[
                    barang["kode_barang"]
                    == kode_barang
                ].iloc[0]

                nama_barang = selected[
                    "nama_barang"
                ]

                satuan = selected[
                    "satuan"
                ]

                st.info(
                    f"Barang: {nama_barang} | "
                    f"Satuan: {satuan}"
                )

            with col2:

                jumlah = st.number_input(
                    "Jumlah",
                    min_value=1,
                    value=1
                )

                supplier_customer = st.text_input(
                    "Supplier / Customer"
                )

                keterangan = st.text_area(
                    "Keterangan"
                )

                foto = st.file_uploader(
                    "Foto Dokumen / Barang",
                    type=[
                        "jpg",
                        "jpeg",
                        "png",
                        "pdf"
                    ]
                )

            submit = st.form_submit_button(
                "💾 Simpan Transaksi",
                type="primary"
            )

            if submit:

                nama_file = ""

                if foto:

                    timestamp = datetime.now().strftime(
                        "%Y%m%d%H%M%S"
                    )

                    nama_file = (
                        timestamp
                        + "_"
                        + foto.name
                    )

                    file_path = os.path.join(
                        UPLOAD_DIR,
                        nama_file
                    )

                    with open(
                        file_path,
                        "wb"
                    ) as f:

                        f.write(
                            foto.getbuffer()
                        )

                tambah_transaksi(
                    nomor_dokumen,
                    str(tanggal),
                    jenis,
                    kode_barang,
                    nama_barang,
                    jumlah,
                    satuan,
                    supplier_customer,
                    keterangan,
                    nama_file
                )

                st.success(
                    f"{jenis} berhasil disimpan."
                )

                st.rerun()


# =========================================================
# RIWAYAT TRANSAKSI
# =========================================================

elif menu == "Riwayat Transaksi":

    st.title("📋 Riwayat Transaksi")

    df = get_transaksi()

    if df.empty:

        st.info(
            "Belum ada transaksi."
        )

    else:

        col1, col2 = st.columns(2)

        with col1:

            jenis_filter = st.selectbox(
                "Jenis",
                [
                    "Semua",
                    "Barang Masuk",
                    "Barang Keluar"
                ]
            )

        with col2:

            keyword = st.text_input(
                "Cari barang / dokumen"
            )

        filtered = df.copy()

        if jenis_filter != "Semua":

            filtered = filtered[
                filtered["jenis"]
                == jenis_filter
            ]

        if keyword:

            filtered = filtered[
                filtered.astype(str)
                .apply(
                    lambda x:
                    x.str.contains(
                        keyword,
                        case=False,
                        na=False
                    )
                )
                .any(axis=1)
            ]

        st.dataframe(
            filtered,
            use_container_width=True,
            hide_index=True
        )


# =========================================================
# EXPORT EXCEL
# =========================================================

elif menu == "Laporan Excel":

    st.title("📊 Laporan Excel")

    st.write(
        "Export seluruh data stok dan transaksi "
        "ke dalam file Excel."
    )

    stok = hitung_stok()
    transaksi = get_transaksi()

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Jenis Barang",
        len(stok)
    )

    col2.metric(
        "Transaksi",
        len(transaksi)
    )

    col3.metric(
        "Total Stok",
        stok["stok"].sum()
        if not stok.empty else 0
    )

    st.divider()

    excel_file = export_excel()

    st.download_button(
        label="📥 Download Laporan Excel",
        data=excel_file,
        file_name=(
            "Laporan_Stok_Gudang_"
            + datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
            + ".xlsx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary"
    )

    st.info(
        "File Excel berisi 3 sheet: "
        "Stok Barang, Riwayat Transaksi, "
        "dan Master Barang."
    )
```
